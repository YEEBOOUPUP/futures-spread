# -*- coding: utf-8 -*-
"""检查用户 Excel 马棕块（国外油脂期货 A-N 列）各列最后日期"""
import openpyxl
from datetime import datetime

PATH = r'C:\Users\10172\OneDrive\Desktop\油脂产业整理-王一波\数据库\油脂油料数据库.xlsx'
wb = openpyxl.load_workbook(PATH, read_only=True, data_only=True)
ws = wb['国外油脂期货']
rows = list(ws.iter_rows(min_row=5, max_row=ws.max_row, max_col=14, values_only=True))
wb.close()
# B=活跃, C-N=01~12 月
labels = ['活跃'] + ['%02d' % i for i in range(1, 13)]
for ci, lb in enumerate(labels):
    last = None
    for r in reversed(rows):
        d, v = r[0], r[ci + 1] if ci + 1 < len(r) else None
        if isinstance(d, datetime) and isinstance(v, (int, float)):
            last = (d.strftime('%Y-%m-%d'), v)
            break
    print('马棕 %-4s: 最后有值 %s' % (lb, ('%s (%.0f)' % last) if last else '-'))
